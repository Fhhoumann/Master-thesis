from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
INPUT_DIR = ROOT / "artifacts" / "appendix_signal_drivers_final_prema"
OUTPUT = ROOT / "reports" / "tables" / "signal_driver_appendix_tables.docx"


NAVY = "27384A"
SECTION_BLUE = "DCE6F1"
GRID = "D9D9D9"
WHITE = "FFFFFF"
TEXT = "111111"


def read_rows(path: Path, limit: int = 10) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        return [row for _, row in zip(range(limit), reader)]


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRID) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text: str, bold: bool = False, italic: bool = False, color: str = TEXT, align=WD_ALIGN_PARAGRAPH.RIGHT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def fmt_number(value: str, decimals: int = 4) -> str:
    return f"{float(value):.{decimals}f}"


def fmt_pct(value: str) -> str:
    return f"{float(value) * 100:.1f}%"


def add_note(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.italic = True
    run.font.color.rgb = RGBColor(80, 80, 80)


def add_table(doc: Document, title: str, category: str, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(4)
    run = heading.add_run(title)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10)

    table = doc.add_table(rows=2, cols=len(headers))
    table.autofit = False
    table.allow_autofit = False

    for idx, width in enumerate(widths):
        for cell in table.columns[idx].cells:
            cell.width = Inches(width)

    header = table.rows[0].cells
    for idx, label in enumerate(headers):
        shade_cell(header[idx], NAVY)
        set_cell_border(header[idx], WHITE)
        set_cell_text(header[idx], label, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.RIGHT)

    category_cells = table.rows[1].cells
    merged = category_cells[0]
    for cell in category_cells[1:]:
        merged = merged.merge(cell)
    shade_cell(merged, SECTION_BLUE)
    set_cell_border(merged)
    set_cell_text(merged, category, bold=True, italic=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.LEFT)

    for values in rows:
        row_cells = table.add_row().cells
        for idx, value in enumerate(values):
            shade_cell(row_cells[idx], WHITE)
            set_cell_border(row_cells[idx])
            set_cell_text(row_cells[idx], value, align=WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.RIGHT)

    doc.add_paragraph()


def main() -> None:
    en = read_rows(INPUT_DIR / "elasticnet_coefficients_aggregated.csv")
    rf = read_rows(INPUT_DIR / "random_forest_importance_aggregated.csv")
    torch = read_rows(INPUT_DIR / "torch_permutation_importance_aggregated.csv")

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    title = doc.add_paragraph()
    title_run = title.add_run("Appendix Signal-Driver Tables")
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_after = Pt(6)

    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run("Final pre-MA specification. Tables report aggregated signal importance across the full out-of-sample period.")
    subtitle_run.font.name = "Arial"
    subtitle_run.font.size = Pt(10)
    subtitle.paragraph_format.space_after = Pt(10)

    add_table(
        doc,
        "Table A.X. Elastic Net signal drivers",
        "Elastic Net coefficients",
        ["Variable", "Mean coef.", "Mean abs. coef.", "Non-zero freq.", "Folds"],
        [[
            row["feature"],
            fmt_number(row["mean_coefficient"]),
            fmt_number(row["mean_abs_coefficient"]),
            fmt_pct(row["nonzero_frequency"]),
            row["fold_count"],
        ] for row in en],
        [3.35, 1.25, 1.45, 1.35, 0.85],
    )
    add_note(doc, "Note: Mean coefficient preserves the sign of the estimated Elastic Net coefficient, while mean absolute coefficient ranks variables by average magnitude. Non-zero frequency reports how often the coefficient is selected across 300 rolling folds.")

    add_table(
        doc,
        "Table A.X. Random Forest signal drivers",
        "Random Forest feature importance",
        ["Variable", "Mean importance", "Std. importance", "Folds"],
        [[
            row["feature"],
            fmt_number(row["mean_importance"]),
            fmt_number(row["std_importance"]),
            row["fold_count"],
        ] for row in rf],
        [3.35, 1.45, 1.45, 0.85],
    )
    add_note(doc, "Note: Random Forest importance is averaged across 300 rolling folds. The values are model-specific importance scores and should be interpreted as relative within-model rankings rather than directly comparable to Elastic Net coefficients or Torch permutation drops.")

    add_table(
        doc,
        "Table A.X. Torch signal drivers",
        "Torch permutation importance",
        ["Variable", "Mean drop", "Std. drop", "Baseline return", "Permuted return", "Obs."],
        [[
            row["feature"],
            fmt_number(row["mean_importance_drop"]),
            fmt_number(row["std_importance_drop"]),
            fmt_number(row["mean_baseline_return"]),
            fmt_number(row["mean_permuted_return"]),
            row["observation_count"],
        ] for row in torch],
        [2.65, 1.05, 1.05, 1.25, 1.25, 0.8],
    )
    add_note(doc, "Note: Torch permutation importance is computed as the reduction in validation-period portfolio return when a feature is permuted. A larger mean drop indicates that disrupting the signal reduces performance more strongly.")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
